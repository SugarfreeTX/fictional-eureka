"""
Bi-weekly savings expectation tracker.

Designed to run from Windows Task Scheduler. Appends a deposit row
only when a payday has arrived (or was missed), so daily/weekly
triggers are safe. Excel is the source of truth for Streamlit later.

Usage:
    python savings_tracker.py              # append any due deposits
    python savings_tracker.py --project 2026-12-31
    python savings_tracker.py --add 250 "bonus / extra transfer"
"""

from __future__ import annotations

import argparse
import logging
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

# ---------------------------------------------------------------------------
# Config — edit these
# ---------------------------------------------------------------------------
EXCEL_PATH = Path(__file__).with_name("savings_ledger.xlsx")
LOG_PATH = Path(__file__).with_name("savings_tracker.log")

# A real payday you already know. All future paydays are +14 days from this.
PAYDAY_ANCHOR = date(2026, 9, 11)  # example Friday; change to your next/last payday
BIWEEKLY_AMOUNT = 750.00          # planned transfer each pay period
STARTING_BALANCE = 5277.00           # cash already in the "account" before first row

# Optional: stop adding scheduled rows after this date (None = no cap)
END_DATE: date | None = None

TX_HEADERS = ["date", "amount_transferred", "running_balance", "source", "notes"]
ASSUMPTION_ROWS = [
    ("payday_anchor", PAYDAY_ANCHOR),
    ("biweekly_amount", BIWEEKLY_AMOUNT),
    ("starting_balance", STARTING_BALANCE),
]


# ---------------------------------------------------------------------------
# Logging
# ---------------------------------------------------------------------------
def setup_logging() -> None:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        handlers=[
            logging.FileHandler(LOG_PATH, encoding="utf-8"),
            logging.StreamHandler(),
        ],
    )


# ---------------------------------------------------------------------------
# Workbook bootstrap
# ---------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF")
INPUT_FONT = Font(name="Calibri", color="0000FF")  # hardcoded inputs
FORMULA_FONT = Font(name="Calibri", color="000000")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
CURRENCY_FORMAT = '"$"#,##0.00'


def _style_header(ws: Worksheet, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def create_workbook(path: Path) -> None:
    wb = Workbook()
    # Remove the default empty sheet so Pylance never sees wb.active as None.
    default_name = wb.sheetnames[0]
    wb.remove(wb[default_name])

    assumptions: Worksheet = wb.create_sheet("Assumptions")
    assumptions["A1"] = "key"
    assumptions["B1"] = "value"
    assumptions["C1"] = "notes"
    _style_header(assumptions, 3)
    assumptions["A2"] = "payday_anchor"
    assumptions["B2"] = PAYDAY_ANCHOR
    assumptions["C2"] = "Known payday; every deposit is +14 days from this"
    assumptions["A3"] = "biweekly_amount"
    assumptions["B3"] = BIWEEKLY_AMOUNT
    assumptions["C3"] = "Planned transfer each pay period"
    assumptions["A4"] = "starting_balance"
    assumptions["B4"] = STARTING_BALANCE
    assumptions["C4"] = "Balance before the first ledger row"
    for row in range(2, 5):
        cell = assumptions.cell(row, 2)
        cell.font = INPUT_FONT
        cell.fill = YELLOW
    assumptions["B2"].number_format = "YYYY-MM-DD"
    assumptions["B3"].number_format = CURRENCY_FORMAT
    assumptions["B4"].number_format = CURRENCY_FORMAT
    assumptions.column_dimensions["A"].width = 22
    assumptions.column_dimensions["B"].width = 16
    assumptions.column_dimensions["C"].width = 52
    assumptions["A6"] = "How to use"
    assumptions["A6"].font = Font(name="Calibri", bold=True)
    assumptions["A7"] = (
        "Yellow/blue cells are inputs. Change the amount here if your "
        "planned transfer changes; new scheduled rows will pick it up. "
        "Do not type over running_balance — it is a formula."
    )
    assumptions.merge_cells("A7:C9")
    assumptions["A7"].alignment = Alignment(wrap_text=True, vertical="top")

    tx: Worksheet = wb.create_sheet("Transactions")
    for col, name in enumerate(TX_HEADERS, start=1):
        tx.cell(1, col, name)
    _style_header(tx, len(TX_HEADERS))
    tx.column_dimensions["A"].width = 14
    tx.column_dimensions["B"].width = 22
    tx.column_dimensions["C"].width = 20
    tx.column_dimensions["D"].width = 14
    tx.column_dimensions["E"].width = 40
    tx.freeze_panes = "A2"
    tx.auto_filter.ref = "A1:E1"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def load_or_create(path: Path):
    if not path.exists():
        create_workbook(path)
        logging.info("Created new ledger at %s", path)
    return load_workbook(path)


def read_assumptions(wb) -> dict:
    ws = wb["Assumptions"]
    out = {}
    for row in range(2, ws.max_row + 1):
        key = ws.cell(row, 1).value
        if not key:
            continue
        out[str(key).strip()] = ws.cell(row, 2).value
    return out


def as_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        return date.fromisoformat(value[:10])
    raise TypeError(f"Cannot parse date from {value!r}")


# ---------------------------------------------------------------------------
# Payday calendar
# ---------------------------------------------------------------------------
def payday_series(anchor: date, until: date) -> list[date]:
    """Inclusive list of bi-weekly dates from anchor through until."""
    if until < anchor:
        return []
    days = (until - anchor).days
    n = days // 14
    return [anchor + timedelta(days=14 * i) for i in range(n + 1)]


def existing_scheduled_dates(ws) -> set[date]:
    dates: set[date] = set()
    for row in range(2, ws.max_row + 1):
        raw = ws.cell(row, 1).value
        source = (ws.cell(row, 4).value or "").strip().lower()
        if raw is None:
            continue
        if source and source != "scheduled":
            continue
        dates.add(as_date(raw))
    return dates


def last_data_row(ws) -> int:
    row = ws.max_row
    while row >= 2 and ws.cell(row, 1).value is None:
        row -= 1
    return row if row >= 2 else 1


# ---------------------------------------------------------------------------
# Writes
# ---------------------------------------------------------------------------
def running_balance_formula(row: int) -> str:
    """Row 2 seeds from Assumptions starting_balance; later rows add prior."""
    if row == 2:
        return "=Assumptions!$B$4+B2"
    return f"=C{row - 1}+B{row}"


def append_row(ws, when: date, amount: float, source: str, notes: str) -> int:
    row = last_data_row(ws) + 1
    ws.cell(row, 1, when).number_format = "YYYY-MM-DD"
    ws.cell(row, 2, float(amount)).number_format = CURRENCY_FORMAT
    ws.cell(row, 2).font = INPUT_FONT
    formula_cell = ws.cell(row, 3, running_balance_formula(row))
    formula_cell.number_format = CURRENCY_FORMAT
    formula_cell.font = FORMULA_FONT
    ws.cell(row, 4, source)
    ws.cell(row, 5, notes)
    return row


def refresh_balance_formulas(ws) -> None:
    """Keep formulas contiguous if someone inserts/deletes a row by hand."""
    for row in range(2, last_data_row(ws) + 1):
        if ws.cell(row, 1).value is None:
            continue
        ws.cell(row, 3, running_balance_formula(row))
        ws.cell(row, 3).number_format = CURRENCY_FORMAT
        ws.cell(row, 3).font = FORMULA_FONT


def apply_table(ws) -> None:
    end_row = max(last_data_row(ws), 2)
    ref = f"A1:E{end_row}"
    # Replace any existing table so the range stays current
    for name in list(ws.tables):
        del ws.tables[name]
    table = Table(displayName="Transactions", ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    ws.auto_filter.ref = ref


# ---------------------------------------------------------------------------
# Core actions
# ---------------------------------------------------------------------------
def add_due_scheduled_rows(wb, today: date | None = None) -> int:
    today = today or date.today()
    assumptions = read_assumptions(wb)
    anchor = as_date(assumptions.get("payday_anchor", PAYDAY_ANCHOR))
    amount = float(assumptions.get("biweekly_amount", BIWEEKLY_AMOUNT))
    until = today if END_DATE is None else min(today, END_DATE)

    ws = wb["Transactions"]
    already = existing_scheduled_dates(ws)
    due = [d for d in payday_series(anchor, until) if d not in already]

    for d in due:
        append_row(ws, d, amount, "scheduled", "bi-weekly planned transfer")
        logging.info("Added scheduled transfer %s on %s", f"{amount:.2f}", d.isoformat())

    if due:
        refresh_balance_formulas(ws)
        apply_table(ws)
    return len(due)


def add_manual_row(wb, amount: float, notes: str, when: date | None = None) -> None:
    when = when or date.today()
    ws = wb["Transactions"]
    append_row(ws, when, amount, "manual", notes)
    refresh_balance_formulas(ws)
    apply_table(ws)
    logging.info("Added manual transfer %s on %s (%s)", f"{amount:.2f}", when.isoformat(), notes)


def expected_balance_on(wb, target: date) -> tuple[float, int]:
    """
    Linear projection: current ledger sum of posted rows through today,
    plus remaining scheduled deposits from tomorrow through target.
    """
    assumptions = read_assumptions(wb)
    anchor = as_date(assumptions.get("payday_anchor", PAYDAY_ANCHOR))
    amount = float(assumptions.get("biweekly_amount", BIWEEKLY_AMOUNT))
    starting = float(assumptions.get("starting_balance", STARTING_BALANCE))

    ws = wb["Transactions"]
    posted = 0.0
    last_posted: date | None = None
    for row in range(2, last_data_row(ws) + 1):
        raw = ws.cell(row, 1).value
        amt = ws.cell(row, 2).value
        if raw is None or amt is None:
            continue
        d = as_date(raw)
        if d <= date.today():
            posted += float(amt)
            last_posted = d if last_posted is None else max(last_posted, d)

    current = starting + posted
    future = [d for d in payday_series(anchor, target) if d > date.today()]
    return current + amount * len(future), len(future)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Bi-weekly savings ledger")
    p.add_argument(
        "--project",
        metavar="YYYY-MM-DD",
        help="Print expected balance on this date (does not write rows past today)",
    )
    p.add_argument(
        "--add",
        nargs=2,
        metavar=("AMOUNT", "NOTE"),
        help="Append a one-off manual transfer today",
    )
    p.add_argument(
        "--as-of",
        metavar="YYYY-MM-DD",
        help="Pretend today is this date (useful for backfill / testing)",
    )
    return p.parse_args()


def main() -> None:
    setup_logging()
    args = parse_args()
    fake_today = date.fromisoformat(args.as_of) if args.as_of else None
    today = fake_today or date.today()

    wb = load_or_create(EXCEL_PATH)
    added = add_due_scheduled_rows(wb, today=today)

    if args.add:
        add_manual_row(wb, float(args.add[0]), args.add[1], when=today)

    wb.save(EXCEL_PATH)
    logging.info("Saved %s (%s scheduled row(s) added)", EXCEL_PATH, added)

    if args.project:
        target = date.fromisoformat(args.project)
        expected, remaining = expected_balance_on(wb, target)
        print(
            f"Expected balance on {target.isoformat()}: "
            f"${expected:,.2f}  ({remaining} future scheduled deposits)"
        )


if __name__ == "__main__":
    main()
